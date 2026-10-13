from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "data/processed/week4_advanced_tuning_results.txt"
OUT = ROOT / "docs/Week4_Advanced_Model_Tuning_Report.docx"

results = RESULTS.read_text(encoding="utf-8")
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)

doc.add_heading("CampusIQ Academic Policy RAG", 0)
doc.add_heading("Week 4 - Advanced Model Tuning Report", level=1)

doc.add_heading("1. Objective", level=2)
doc.add_paragraph(
    "Compare a basic text-classification model with advanced features "
    "and tuned machine-learning models."
)

doc.add_heading("2. Dataset", level=2)
doc.add_paragraph(
    "There were 62 manually labeled examples. The academic_calendar "
    "category, with one example, was excluded. The final dataset contained "
    "61 examples: 45 training examples and 16 test examples. Labels are "
    "provisional and have not been independently verified."
)

doc.add_heading("3. Feature Engineering", level=2)
for item in [
    "Word TF-IDF using unigrams and bigrams.",
    "Character-level TF-IDF using character n-grams.",
    "Numeric features: character count, word count, sentence count, "
    "uppercase ratio, digit ratio, and policy keyword counts.",
    "StandardScaler applied to numerical features.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("4. Models and Tuning", level=2)
doc.add_paragraph(
    "Models: majority-class DummyClassifier, word-only TF-IDF Logistic "
    "Regression, tuned Logistic Regression, and tuned Linear SVM. "
    "GridSearchCV tested C values 0.1, 1.0, and 10.0 using three-fold "
    "stratified cross-validation and macro F1 scoring."
)

doc.add_heading("5. Results", level=2)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for cell, value in zip(
    table.rows[0].cells,
    ["Model", "Accuracy", "Macro F1", "Weighted F1"]
):
    cell.text = value

rows = [
    ("Majority baseline", "37.50%", "0.109", "0.205"),
    ("Word TF-IDF Logistic Regression", "87.50%", "0.876", "0.869"),
    ("Tuned Logistic Regression + interaction features", "62.50%", "0.639", "0.610"),
    ("Tuned Linear SVM + interaction features", "62.50%", "0.639", "0.610"),

]
for row in rows:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value

 
doc.add_paragraph(
    "After adding interaction features, the best cross-validation "
    "macro F1 was 0.615 for Logistic Regression (C=0.1) and "
    "0.595 for Linear SVM (C=0.1)."
)


doc.add_heading("6. Interpretation", level=2)
 
doc.add_paragraph(
    "The word-only TF-IDF Logistic Regression baseline performed "
    "best on this test split, achieving 87.5% accuracy and a macro "
    "F1 of 0.876. After adding interaction features and tuning, both "
    "advanced models achieved 62.5% accuracy and a macro F1 of 0.639. "
    "Therefore, the added features and tuning did not improve the "
    "held-out test result. The dataset is small, so these findings "
    "are preliminary."
)


doc.add_heading("7. Limitations", level=2)
for item in [
    "The dataset and test set are small.",
    "Labels are provisional and not independently verified.",
    "One category was excluded because it had only one example.",
    "Results from one split do not prove real-world generalization.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("8. Conclusion and Future Work", level=2)
doc.add_paragraph(
    "The simpler word-only TF-IDF Logistic Regression model was the "
    "test-split winner. Future work includes reviewing labels, collecting "
    "more examples, using additional validation splits, and separately "
    "evaluating RAG retrieval and answer grounding. The model is not "
    "yet production-ready."
)

doc.add_heading("9. Experiment Output", level=2)
doc.add_paragraph(results)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("Report created successfully:", OUT)
